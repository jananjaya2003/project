import streamlit as st
import pandas as pd
import openpyxl
import os
import plotly.express as px

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "Supplier Registration status.xlsx")
HEADER_COLUMNS = [
    'No',
    'Institute',
    'Payments closing date',
    'Applications Closing Date',
    'Application Mode',
    'Method of payment',
    'Registration Fee',
    'Other Charges',
    'Application sent Date',
    'Remarks',
    'Status'
]

st.set_page_config(
    page_title="Supplier Registration Dashboard",
    layout="wide"
)

AUTH_USERNAME = "admin"
AUTH_PASSWORD = "admin123"

if "authenticated" not in st.session_state:
    st.session_state.authenticated = False


def login_screen():
    st.markdown(
        """
        <div style="max-width: 460px; margin: 4rem auto 2rem auto; padding: 2rem; border-radius: 18px; background: rgba(255,255,255,0.8); box-shadow: 0 12px 30px rgba(15,23,42,0.08); border: 1px solid rgba(148,163,184,0.25);">
            <h2 style="text-align:center; margin-bottom: 1.5rem;">🔐 Login</h2>
        </div>
        """,
        unsafe_allow_html=True,
    )

    with st.form("login_form", clear_on_submit=False):
        username = st.text_input("Username")
        password = st.text_input("Password", type="password")
        submit = st.form_submit_button("Login", use_container_width=True)

        if submit:
            if username == AUTH_USERNAME and password == AUTH_PASSWORD:
                st.session_state.authenticated = True
                st.rerun()
            else:
                st.error("Invalid username or password.")


if not st.session_state.authenticated:
    login_screen()
    st.stop()

st.markdown(
    """
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.3/dist/css/bootstrap.min.css" rel="stylesheet">
    <style>
        .stApp {
            background: linear-gradient(135deg, #f4f7fb 0%, #eef6ff 45%, #f6f1ff 100%);
        }
        .block-container {
            padding-top: 2rem;
            padding-bottom: 2rem;
        }
        [data-testid="stMetric"] {
            background: rgba(255, 255, 255, 0.9);
            border-radius: 20px;
            border: 1px solid rgba(13, 110, 253, 0.08);
            box-shadow: 0 10px 30px rgba(15, 23, 42, 0.08);
            padding: 1.2rem 1rem !important;
            margin-top: 0.5rem;
        }
        [data-testid="stMetricLabel"] {
            color: #334155;
            font-weight: 600;
        }
        [data-testid="stMetricValue"] {
            font-size: 1.7rem !important;
            font-weight: 700;
            color: #0f172a;
        }
        [data-testid="stMetricDelta"] {
            font-size: 0.85rem;
            font-weight: 600;
        }
        .dashboard-header {
            background: linear-gradient(135deg, rgba(13,110,253,0.12), rgba(111,66,193,0.12));
            border: 1px solid rgba(13,110,253,0.18);
            border-radius: 24px;
            padding: 1.5rem 1.75rem;
            margin-bottom: 1.5rem;
            box-shadow: 0 15px 35px rgba(15, 23, 42, 0.08);
        }
        .dashboard-header .badge {
            letter-spacing: 0.08em;
            text-transform: uppercase;
            font-size: 0.75rem;
        }
        .dashboard-header h1 {
            font-size: clamp(2rem, 3vw, 2.9rem);
            font-weight: 800;
            margin: 0;
            color: #0f172a;
        }
        .dashboard-header .status-pill {
            background: rgba(25, 135, 84, 0.1);
            color: #0f5132;
            border: 1px solid rgba(25, 135, 84, 0.2);
            border-radius: 999px;
            padding: 0.55rem 1rem;
            font-weight: 700;
        }
        .stDataFrame, .stPlotlyChart, .stTabs [role="tablist"] {
            border-radius: 18px;
        }
        .stTabs [role="tablist"] {
            gap: 0.5rem;
            background: rgba(255,255,255,0.5);
            padding: 0.25rem;
        }
        .stTabs button {
            border-radius: 12px !important;
            font-weight: 600;
        }
        .stButton > button {
            border-radius: 12px !important;
            font-weight: 600 !important;
            transition: all 0.2s ease;
        }
        .stButton > button:hover {
            transform: translateY(-1px);
            box-shadow: 0 8px 18px rgba(13, 110, 253, 0.18);
        }
    </style>
    """,
    unsafe_allow_html=True,
)


def ensure_excel_structure():
    try:
        if not os.path.exists(EXCEL_FILE):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            for col_idx, column_name in enumerate(HEADER_COLUMNS, start=1):
                ws.cell(row=2, column=col_idx, value=column_name)
            wb.save(EXCEL_FILE)
            wb.close()
            return

        wb = openpyxl.load_workbook(EXCEL_FILE)

        if "Sheet1" not in wb.sheetnames:
            ws = wb.active
            ws.title = "Sheet1"

        ws = wb["Sheet1"]
        if ws.max_row == 0 or ws.max_row == 1:
            ws.append([])

        if ws.max_row == 1 or all(ws.cell(row=1, column=i).value is None for i in range(1, len(HEADER_COLUMNS) + 1)):
            for col_idx, column_name in enumerate(HEADER_COLUMNS, start=1):
                ws.cell(row=1, column=col_idx, value=column_name)

        if ws.max_row < 2:
            for col_idx, column_name in enumerate(HEADER_COLUMNS, start=1):
                ws.cell(row=2, column=col_idx, value=column_name)

        if "Settlement 01" not in wb.sheetnames:
            wb.create_sheet("Settlement 01")
        if "Settlement 02" not in wb.sheetnames:
            wb.create_sheet("Settlement 02")

        wb.save(EXCEL_FILE)
        wb.close()
    except Exception as exc:
        st.error(f"Unable to initialize the Excel file: {exc}")
        raise


def load_data():
    ensure_excel_structure()

    try:
        df_main = pd.read_excel(
            EXCEL_FILE,
            sheet_name="Sheet1",
            skiprows=1
        )
    except Exception as e:
        st.error(f"Error loading Sheet1: {e}")
        return None, None, None

    cols = HEADER_COLUMNS

    df_main = df_main.iloc[:, :11]

    while len(df_main.columns) < 11:
        df_main[f'Col_{len(df_main.columns)}'] = None

    df_main.columns = cols
    df_main = df_main.dropna(subset=['Institute'])

    df_main['Status'] = (
        df_main['Status']
        .fillna('Pending')
        .astype(str)
        .str.strip()
        .str.title()
    )

    try:
        df_set1 = pd.read_excel(EXCEL_FILE, sheet_name='Settlement 01')
    except:
        df_set1 = None

    try:
        df_set2 = pd.read_excel(EXCEL_FILE, sheet_name='Settlement 02')
    except:
        df_set2 = None

    return df_main, df_set1, df_set2


st.markdown(
    """
    <div class="dashboard-header">
        <div class="d-flex justify-content-between align-items-center flex-wrap gap-3">
            <div>
                <span class="badge bg-primary-subtle text-primary-emphasis rounded-pill mb-3">Supplier Portal</span>
                <h1>📊 Supplier Registration Tracking Dashboard</h1>
            </div>
            <div class="d-flex align-items-center gap-2">
                <span class="status-pill">Live Tracking</span>
            </div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)

col_title, col_logout = st.columns([5, 1])
with col_logout:
    if st.button("Logout", key="logout_btn", use_container_width=True):
        st.session_state.authenticated = False
        st.rerun()

df_main, df_set1, df_set2 = load_data()

if df_main is not None:

    st.markdown("### 📈 Overview Analytics")

    m1, m2, m3, m4, m5 = st.columns(5)

    total_inst = len(df_main)

    completed_count = len(
        df_main[df_main['Status'] == 'Complete']
    )

    pending_count = len(
        df_main[df_main['Status'] == 'Pending']
    )

    total_fee = pd.to_numeric(
        df_main['Registration Fee'],
        errors='coerce'
    ).sum()

    m1.metric("Total Institutes", total_inst)
    m2.metric("✅ Completed", completed_count)
    m3.metric("⏳ Pending", pending_count)
    m4.metric("Total Reg Fee (LKR)", f"{total_fee:,.2f}")

    if df_set1 is not None and not df_set1.empty:
        spent_1 = df_set1.iloc[2, 2] if len(df_set1) > 2 else 0
        balance_1 = df_set1.iloc[3, 2] if len(df_set1) > 3 else 0

        spent_1 = pd.to_numeric(spent_1, errors="coerce")
        balance_1 = pd.to_numeric(balance_1, errors="coerce")

        spent_1 = 0 if pd.isna(spent_1) else spent_1
        balance_1 = 0 if pd.isna(balance_1) else balance_1

        m5.metric(
            "Settlement 01 Spent",
            f"{spent_1:,.2f}",
            delta=f"Bal: {balance_1:,.2f}"
        )

    st.markdown("---")

    search = st.text_input(
        "🔍 Search Institute",
        placeholder="Enter institute name"
    )

    df_display = df_main.copy()
    if search:
        df_display = df_display[
            df_display['Institute']
            .astype(str)
            .str.contains(search, case=False, na=False)
        ]

    st.subheader("📋 Current Registration Records")

    def color_status(val):
        if str(val).strip().lower() == 'complete':
            return (
                'background-color:#d4edda;'
                'color:#155724;'
                'font-weight:bold;'
            )
        return (
            'background-color:#f8d7da;'
            'color:#721c24;'
            'font-weight:bold;'
        )

    if hasattr(df_display.style, 'map'):
        styled_df = df_display.style.map(
            color_status,
            subset=['Status']
        )
    else:
        styled_df = df_display.style.applymap(
            color_status,
            subset=['Status']
        )

    st.dataframe(
        styled_df,
        width="stretch"
    )

    st.markdown("---")
    st.subheader("📊 Completion Status Breakdown")

    status_counts = (
        df_display['Status']
        .value_counts()
        .reset_index()
    )

    status_counts.columns = [
        'Status',
        'Count'
    ]

    if not status_counts.empty:
        fig_status = px.pie(
            status_counts,
            values="Count",
            names="Status",
            color="Status",
            color_discrete_map={
                "Complete": "#28a745",
                "Pending": "#dc3545"
            },
            hole=0.4,
            title="Records Completion Status"
        )

        st.plotly_chart(
            fig_status,
            width="stretch"
        )

    st.markdown("---")

    tab_add, tab_edit, tab_delete = st.tabs([
        "➕ Add New Record",
        "✏️ Edit Record by No",
        "🗑️ Delete Record by No"
    ])

    # ==========================
    # ADD RECORD
    # ==========================
    with tab_add:
        st.subheader("Add New Record")

        with st.form("add_record_form", clear_on_submit=True):
            numeric_nos = pd.to_numeric(
                df_main["No"],
                errors="coerce"
            ).dropna()

            next_no = 1 if numeric_nos.empty else int(numeric_nos.max()) + 1

            col1, col2, col3 = st.columns(3)

            with col1:
                no_input = st.number_input("Record No", value=next_no, step=1)
                institute = st.text_input("Institute Name *")
                pay_close_date = st.date_input(
                    "Payments Closing Date",
                    value=None,
                    min_value=None,
                    max_value=None
                )
                status_input = st.selectbox("Status", ["Pending", "Complete"])

            with col2:
                app_close_date = st.date_input(
                    "Applications Closing Date",
                    value=None,
                    min_value=None,
                    max_value=None
                )
                app_mode = st.selectbox(
                    "Application Mode",
                    ["Self", "Hospital", "Post", "Online", "Other"]
                )
                method_payment = st.selectbox(
                    "Method of Payment",
                    ["Shroff", "BOC", "Money Order", "Bank Deposit", "Online", "Other"]
                )

            with col3:
                reg_fee = st.number_input("Registration Fee", min_value=0.0)
                other_charges = st.number_input("Other Charges", min_value=0.0)
                app_sent_date = st.date_input(
                    "Application Sent Date",
                    value=None,
                    min_value=None,
                    max_value=None
                )
                remarks = st.text_input("Remarks")

            submitted = st.form_submit_button("💾 Save Record")

            if submitted:
                if not institute.strip():
                    st.error("Institute name is required.")
                else:
                    existing_nos = pd.to_numeric(
                        df_main["No"],
                        errors="coerce"
                    ).tolist()

                    if no_input in existing_nos:
                        st.error("Record number already exists.")
                    else:
                        try:
                            wb = openpyxl.load_workbook(EXCEL_FILE)
                            ws = wb["Sheet1"]

                            row_idx = 3
                            while ws.cell(row=row_idx, column=2).value is not None:
                                row_idx += 1

                            values = [
                                no_input,
                                institute,
                                pay_close_date.strftime('%Y-%m-%d') if pay_close_date is not None else '',
                                app_close_date.strftime('%Y-%m-%d') if app_close_date is not None else '',
                                app_mode,
                                method_payment,
                                reg_fee,
                                other_charges if other_charges > 0 else None,
                                app_sent_date.strftime('%Y-%m-%d') if app_sent_date is not None else '',
                                remarks,
                                status_input
                            ]

                            for col, value in enumerate(values, start=1):
                                ws.cell(row=row_idx, column=col, value=value)

                            wb.save(EXCEL_FILE)
                            st.success(f"Record No. {no_input} added successfully.")
                            st.rerun()

                        except Exception as e:
                            st.error(f"Save failed: {e}")

    # ==========================
    # EDIT RECORD BY NO
    # ==========================
    with tab_edit:
        st.subheader("Edit Record by Number (No)")

        valid_nos = (
            df_main["No"]
            .dropna()
            .astype(str)
            .tolist()
        )

        if not valid_nos:
            st.info("No records available to edit.")
        else:
            selected_no = st.selectbox("Select Record Number (No) to Edit", valid_nos, key="edit_no_select")

            try:
                wb = openpyxl.load_workbook(EXCEL_FILE)
                ws = wb["Sheet1"]

                target_row = None
                for r in range(3, ws.max_row + 1):
                    val = ws.cell(row=r, column=1).value
                    if val is not None:
                        try:
                            if int(val) == int(selected_no):
                                target_row = r
                                break
                        except:
                            pass

                if target_row:
                    curr_inst = ws.cell(row=target_row, column=2).value or ""
                    curr_pay_date = ws.cell(row=target_row, column=3).value or ""
                    curr_app_date = ws.cell(row=target_row, column=4).value or ""
                    curr_mode = ws.cell(row=target_row, column=5).value or "Self"
                    curr_method = ws.cell(row=target_row, column=6).value or "Shroff"
                    curr_fee = float(ws.cell(row=target_row, column=7).value or 0)
                    curr_charges = ws.cell(row=target_row, column=8).value
                    curr_charges = float(curr_charges) if curr_charges is not None else 0.0
                    curr_sent = ws.cell(row=target_row, column=9).value or ""
                    curr_remarks = ws.cell(row=target_row, column=10).value or ""
                    curr_status = str(ws.cell(row=target_row, column=11).value or "Pending").title()

                    with st.form("edit_form"):
                        col1, col2, col3 = st.columns(3)
                        default_pay_close = pd.to_datetime(curr_pay_date, errors='coerce').to_pydatetime() if str(curr_pay_date).strip() else None
                        default_app_close = pd.to_datetime(curr_app_date, errors='coerce').to_pydatetime() if str(curr_app_date).strip() else None
                        default_sent_date = pd.to_datetime(curr_sent, errors='coerce').to_pydatetime() if str(curr_sent).strip() else None

                        with col1:
                            st.markdown(f"**Editing Record No: `{selected_no}`**")
                            edit_inst = st.text_input("Institute Name *", value=str(curr_inst))
                            edit_pay_close = st.date_input("Payments Closing Date", value=default_pay_close)
                            status_opts = ["Pending", "Complete"]
                            edit_status = st.selectbox(
                                "Status", 
                                status_opts, 
                                index=status_opts.index(curr_status) if curr_status in status_opts else 0
                            )

                        with col2:
                            edit_app_close = st.date_input("Applications Closing Date", value=default_app_close)
                            mode_opts = ["Self", "Hospital", "Post", "Online", "Other"]
                            edit_mode = st.selectbox(
                                "Application Mode", 
                                mode_opts, 
                                index=mode_opts.index(curr_mode) if curr_mode in mode_opts else 0
                            )
                            method_opts = ["Shroff", "BOC", "Money Order", "Bank Deposit", "Online", "Other"]
                            edit_method = st.selectbox(
                                "Method of Payment", 
                                method_opts, 
                                index=method_opts.index(curr_method) if curr_method in method_opts else 0
                            )

                        with col3:
                            edit_fee = st.number_input("Registration Fee", min_value=0.0, value=curr_fee)
                            edit_charges = st.number_input("Other Charges", min_value=0.0, value=curr_charges)
                            edit_sent = st.date_input("Application Sent Date", value=default_sent_date)
                            edit_remarks = st.text_input("Remarks", value=str(curr_remarks))

                        update_btn = st.form_submit_button("💾 Update Record")

                        if update_btn:
                            ws.cell(row=target_row, column=2, value=edit_inst)
                            ws.cell(row=target_row, column=3, value=edit_pay_close.strftime('%Y-%m-%d') if edit_pay_close is not None else '')
                            ws.cell(row=target_row, column=4, value=edit_app_close.strftime('%Y-%m-%d') if edit_app_close is not None else '')
                            ws.cell(row=target_row, column=5, value=edit_mode)
                            ws.cell(row=target_row, column=6, value=edit_method)
                            ws.cell(row=target_row, column=7, value=edit_fee)
                            ws.cell(row=target_row, column=8, value=edit_charges if edit_charges > 0 else None)
                            ws.cell(row=target_row, column=9, value=edit_sent.strftime('%Y-%m-%d') if edit_sent is not None else '')
                            ws.cell(row=target_row, column=10, value=edit_remarks)
                            ws.cell(row=target_row, column=11, value=edit_status)

                            wb.save(EXCEL_FILE)
                            st.success(f"Record No. {selected_no} updated successfully.")
                            st.rerun()

            except Exception as e:
                st.error(f"Edit failed: {e}")

    # ==========================
    # DELETE RECORD BY NO
    # ==========================
    with tab_delete:
        st.subheader("Delete Record by Number (No)")

        valid_nos_del = (
            df_main["No"]
            .dropna()
            .astype(str)
            .tolist()
        )

        if not valid_nos_del:
            st.info("No records available to delete.")
        else:
            del_no = st.selectbox("Select Record Number (No) to Delete", valid_nos_del, key="del_no_select")

            record_to_del = df_main[df_main['No'].astype(str) == str(del_no)]
            if not record_to_del.empty:
                st.write(f"**Previewing Record No. {del_no} to Delete:**")
                st.dataframe(record_to_del, width="stretch")

            confirm_delete = st.checkbox("I confirm I want to permanently delete this record.", key="del_confirm")

            if st.button("🗑️ Delete Record", type="primary"):
                if not confirm_delete:
                    st.warning("Please check the confirmation box before deleting the record.")
                else:
                    try:
                        wb = openpyxl.load_workbook(EXCEL_FILE)
                        ws = wb["Sheet1"]

                        target_row = None
                        for r in range(3, ws.max_row + 1):
                            val = ws.cell(row=r, column=1).value
                            try:
                                if int(val) == int(del_no):
                                    target_row = r
                                    break
                            except:
                                pass

                        if target_row:
                            ws.delete_rows(target_row)
                            wb.save(EXCEL_FILE)
                            st.success(f"Record No. {del_no} deleted successfully.")
                            st.rerun()
                        else:
                            st.error("Record number not found in sheet.")

                    except Exception as e:
                        st.error(f"Delete failed: {e}")